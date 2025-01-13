from docx import Document
import os
from datetime import datetime
# pip install python-docx

# documento_Word = Document()

# documento_Word.add_heading('Trabalhando com Documento Word.',level=0)
# documento_Word.add_heading(f'{os.linesep}',level=1)
# documento_Word.add_paragraph('olá... isso é um prazer trabalhar com você...')
# documento_Word.save('demo.docx')

data_ano = datetime.now().strftime('%d-%m-%Y')
segundos = datetime.now().strftime('-%S')

print(data_ano)
print(segundos)